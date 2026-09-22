#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-proposal-docx.py
======================
Generate the client-facing Word proposal that twins the print-ready HTML.

    docs/proposal/API-Integration-and-Tally-Export.html   <- source of truth
    docs/proposal/API-Integration-and-Tally-Export.docx   <- this script's output

The HTML is the single source of truth for wording.  Only layout is adapted
(Word has no CSS grid, so 2-up grids become 2-column tables with a spacer
column, and the "chip" becomes a shaded run).

Usage:
    python3 scripts/build-proposal-docx.py

Local only: no network, no installs.  Requires python-docx (already present).
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# geometry
# ---------------------------------------------------------------------------
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(REPO_ROOT, "docs", "proposal", "assets")
FIG_FLOW = os.path.join(ASSETS, "diagram-flow.png")
FIG_TIMELINE = os.path.join(ASSETS, "diagram-timeline.png")
OUT_PATH = os.path.join(
    REPO_ROOT, "docs", "proposal", "API-Integration-and-Tally-Export.docx"
)

PAGE_W, PAGE_H = 21.0, 29.7  # A4, cm
MARGIN_TB, MARGIN_LR = 1.4, 1.5  # cm
CONTENT_W = round(PAGE_W - 2 * MARGIN_LR, 2)  # 18.0 cm
IMAGE_W = 17.9  # cm, full content width (requirement)

# ---------------------------------------------------------------------------
# typeface + palette
# (brand tokens from the task + the extra tones the HTML itself uses)
# ---------------------------------------------------------------------------
FONT = "Calibri"

NAVY = "0F2A44"
GREEN = "1E5C38"
BODY = "1E293B"
MUTED = "94A3B8"
BORDER = "DFE5EA"
SOFT = "F5F8FA"
GREEN_TINT = "EDF3EC"
AMBER = "8A5A00"
AMBER_TINT = "FCF3E0"
WHITE = "FFFFFF"
# tones lifted verbatim from the HTML palette / day cards
SUB = "5A6B78"          # --sub
INK = "12212E"          # --ink
DAY_TAG = "8FC7A8"      # DAY n on navy
DAY_TAG_AMBER = "F5D9A0"  # DAY 0 on amber
AMBER_BODY = "6B4A0A"   # amber body copy
AMBER_EDGE = "E8D9B8"   # DAY 0 card border
COST_LABEL = "3C5A48"   # cost table labels

# ---------------------------------------------------------------------------
# type scale   (HTML body is 8.4pt; Word body is 9pt -> ~1.07 scale)
# ---------------------------------------------------------------------------
SZ_BODY = 9
SZ_CARD = 8.6
SZ_TABLE = 8.2
SZ_SMALL = 8
SZ_BULLET = 7.6
SZ_MICRO = 7.2
SZ_KICKER = 7.5
SZ_H1 = 17
SZ_H1_P2 = 15.5
SZ_H2 = 10.5

# ---------------------------------------------------------------------------
# OOXML child order (a wrong order makes Word reject the file)
# ---------------------------------------------------------------------------
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
    "strike", "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
    "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
    "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
    "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
    "specVanish", "oMath",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription", "tblPrChange",
]
TCBORDERS_ORDER = [
    "top", "start", "left", "bottom", "end", "right", "insideH", "insideV",
]
PBDR_ORDER = ["top", "left", "bottom", "right", "between", "bar"]
TCMAR_ORDER = ["top", "start", "left", "bottom", "end", "right"]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]

TWIPS_PER_CM = 567.0


def _insert_ordered(parent, element, order):
    """Insert `element` into `parent` respecting the schema child order."""
    tag = element.tag.split("}")[-1]
    idx = order.index(tag) if tag in order else len(order)
    for child in parent:
        ctag = child.tag.split("}")[-1]
        cidx = order.index(ctag) if ctag in order else len(order)
        if cidx > idx:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


def _cm2dxa(cm):
    return str(int(round(cm * TWIPS_PER_CM)))


# ---------------------------------------------------------------------------
# low-level XML helpers
# ---------------------------------------------------------------------------
def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _insert_ordered(tcPr, shd, TCPR_ORDER)


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _insert_ordered(pPr, shd, PPR_ORDER)


def set_cell_margins(cell, top=None, left=None, bottom=None, right=None):
    """Cell padding in cm (per-cell override of the table default)."""
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.find(qn("w:tcMar"))
    if mar is None:
        mar = _insert_ordered(tcPr, OxmlElement("w:tcMar"), TCPR_ORDER)
    for name, value in (("top", top), ("left", left), ("bottom", bottom),
                        ("right", right)):
        if value is None:
            continue
        el = mar.find(qn("w:" + name))
        if el is None:
            el = _insert_ordered(mar, OxmlElement("w:" + name), TCMAR_ORDER)
        el.set(qn("w:w"), _cm2dxa(value))
        el.set(qn("w:type"), "dxa")


def set_cell_borders(cell, sides):
    """sides: {'top': (color, sz_eighths_pt, style), ...}"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = _insert_ordered(tcPr, OxmlElement("w:tcBorders"), TCPR_ORDER)
    for name in ("top", "left", "bottom", "right"):
        spec = sides.get(name)
        if not spec:
            continue
        color, sz, style = spec
        el = OxmlElement("w:{}".format(name))
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        _insert_ordered(borders, el, TCBORDERS_ORDER)


def set_paragraph_borders(paragraph, sides):
    """sides: {'left': (color, sz_eighths_pt, style, space_pt), ...}"""
    pPr = paragraph._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _insert_ordered(pPr, OxmlElement("w:pBdr"), PPR_ORDER)
    for name in ("top", "left", "bottom", "right"):
        spec = sides.get(name)
        if not spec:
            continue
        color, sz, style, space = spec
        el = OxmlElement("w:{}".format(name))
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        _insert_ordered(bdr, el, PBDR_ORDER)


def set_cell_vmerge(cell, val="continue"):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:vMerge")
    el.set(qn("w:val"), val)
    _insert_ordered(tcPr, el, TCPR_ORDER)


def clear_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:{}".format(side))
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    _insert_ordered(tblPr, borders, TBLPR_ORDER)


def set_table_width(table, cm):
    tblPr = table._tbl.tblPr
    el = OxmlElement("w:tblW")
    el.set(qn("w:w"), _cm2dxa(cm))
    el.set(qn("w:type"), "dxa")
    _insert_ordered(tblPr, el, TBLPR_ORDER)


def set_table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    el = OxmlElement("w:tblLayout")
    el.set(qn("w:type"), "fixed")
    _insert_ordered(tblPr, el, TBLPR_ORDER)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    _insert_ordered(trPr, OxmlElement("w:cantSplit"), TRPR_ORDER)


def set_col_widths(table, widths_cm):
    """Set the grid column widths AND cell.width on every cell."""
    for idx, width in enumerate(widths_cm):
        col = table.columns[idx]
        col.width = Cm(width)
        for cell in col.cells:
            cell.width = Cm(width)


def add_table(container, rows, cols, widths_cm, borders=False):
    table = container.add_table(rows=rows, cols=cols)
    table.autofit = False
    set_table_fixed_layout(table)
    set_table_width(table, sum(widths_cm))
    set_col_widths(table, widths_cm)
    if not borders:
        clear_table_borders(table)
    return table


def style_run(run, size=SZ_BODY, bold=False, italic=False, color=BODY,
              font=FONT, shade=None, spacing=None, caps=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is not None:  # set by run.font.name above
        rFonts.set(qn("w:eastAsia"), font)
        rFonts.set(qn("w:cs"), font)
    if caps:
        _insert_ordered(rPr, OxmlElement("w:caps"), RPR_ORDER)
    if spacing is not None:  # letter-spacing, in twentieths of a point
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(spacing)))
        _insert_ordered(rPr, el, RPR_ORDER)
    if shade is not None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)
        _insert_ordered(rPr, shd, RPR_ORDER)
    return run


def rich(paragraph, parts, size=SZ_BODY, bold=False, italic=False, color=BODY,
         font=FONT, spacing=None, caps=False, shade=None):
    """Add runs to a paragraph. `parts` is a str or [(text, {overrides}), ...]."""
    if isinstance(parts, str):
        parts = [(parts, {})]
    for text, over in parts:
        style_run(
            paragraph.add_run(text),
            size=over.get("size", size),
            bold=over.get("bold", bold),
            italic=over.get("italic", italic),
            color=over.get("color", color),
            font=over.get("font", font),
            shade=over.get("shade", shade),
            spacing=over.get("spacing", spacing),
            caps=over.get("caps", caps),
        )


def add_page_field(paragraph, size=SZ_MICRO, color=MUTED):
    """A live PAGE field (Word/LibreOffice recompute it on layout)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, cached, end):
        run._r.append(el)
    style_run(run, size=size, color=color)
    return run


# ---------------------------------------------------------------------------
# paragraph factory
# ---------------------------------------------------------------------------
def fmt_para(paragraph, style=None, size=SZ_BODY, align=None, space_before=0,
             space_after=4, line_spacing=1.12, left_indent=None,
             right_indent=None, first_line_indent=None, keep_with_next=False,
             keep_together=False):
    if style:
        paragraph.style = style
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    if right_indent is not None:
        pf.right_indent = Cm(right_indent)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together
    return paragraph


class Flow(object):
    """Paragraph factory.

    Inside a table cell it reuses the cell's mandatory leading empty paragraph;
    at body level (reuse_first=False) it always appends, so the flow order is
    exactly the order the code calls it.
    """

    def __init__(self, container, reuse_first=True):
        self._c = container
        self._first = True
        self._reuse = reuse_first

    def para(self, **fmt):
        paragraph = None
        if self._first:
            self._first = False
            existing = None
            if self._reuse:
                try:
                    existing = self._c.paragraphs[0]
                except (IndexError, AttributeError):
                    existing = None
            if existing is not None and not existing.text and not existing.runs:
                paragraph = existing
        if paragraph is None:
            paragraph = self._c.add_paragraph()
        return fmt_para(paragraph, **fmt)

    def bullet(self, parts, size=SZ_BULLET, color=BODY, space_after=1.2,
               indent=0.34, keep_with_next=False):
        paragraph = self.para(style="List Bullet", space_after=space_after,
                              space_before=0, line_spacing=1.06,
                              left_indent=indent, first_line_indent=-indent,
                              keep_with_next=keep_with_next)
        rich(paragraph, parts, size=size, color=color)
        return paragraph


# ===========================================================================
# CONTENT  — transcribed from the HTML (source of truth)
# ===========================================================================
LEDE_P1 = ("How e-Invoicing (IRN + E-Way Bill) and Tally accounting will run "
           "automatically from inside the CRM — and a five-day plan to set it up.")
META_P1 = [("Prepared for:", "Prokon"),
           ("Scope:", "3 branches · 1 Tally company"),
           ("Delivery:", "5 working days"),
           ("Version:", "1.0")]

CARD_A = {
    "tag": "A · COMPLIANCE",
    "title": "GSP API — E-Invoice (IRN) & E-Way Bill",
    "body": ("The CRM itself fetches the legally valid IRN, acknowledgement "
             "number, signed QR code and E-Way Bill number from the Government "
             "portal, through our GSP partner. No separate portal login, no "
             "copy-paste."),
    "chip": "₹29,500 / year · 4,000 API calls",
}
CARD_B = {
    "tag": "B · ACCOUNTING",
    "title": "Tally Export — books stay in TallyPrime",
    "body": ("Every CRM invoice is converted into TallyPrime vouchers with "
             "matching ledgers, GST rate buckets and invoice references. One "
             "file, imported in one go — Sales, GST and party ledgers update "
             "together."),
    "chip": "₹0 — Tally's import gateway is built in",
}
RULE = [("The rule both follow:", {"bold": True, "color": GREEN}),
        (" the CRM is the only place a bill is created. Nothing is re-typed in "
         "the Government portal or in Tally — so the same figures reach the tax "
         "department and your books.", {})]

# label order as specified for the Word legend table
LEGEND = [
    ("1", "Invoice created in the CRM — any of the three branches, exactly as today."),
    ("2A", "Validation, then a secure server-to-server call to the GSP API. Credentials never reach the browser."),
    ("3A", "Government returns IRN, acknowledgement number, signed QR and E-Way Bill number; the invoice is marked Complete."),
    ("2B", "Export engine converts the invoice into Tally vouchers and checks every figure first."),
    ("3B", "One XML file imported into a Tally trial company first, then into your live company."),
    ("★", "Both paths start from the same stored invoice — that is what keeps taxes and books in agreement."),
]

CHANGES_HEAD = ["Task", "Today — manual", "After integration — automatic"]
CHANGES = [
    ("IRN / e-Invoice",
     "Download JSON → upload to GST portal → copy the response → paste it back into the CRM",
     "One click in the CRM. IRN, acknowledgement number and QR return and save themselves."),
    ("E-Way Bill",
     "A second manual round trip on a separate portal",
     "Raised immediately after the IRN, with transport details prefilled and validity date stored."),
    ("Tally entry",
     "Accounts re-types every invoice into TallyPrime by hand",
     "Pick a date range and branch, export one XML file, import it — all vouchers post at once."),
    ("Errors & audit",
     "Mistakes surface late, usually at GST filing time",
     "Balance and GST-rate checks block a bad voucher; every API call is logged and traceable."),
    ("Reference integrity",
     "A wrong or invented bill number cannot be detected",
     "The system refuses to save or export a statutory reference that was not genuinely issued."),
]

LEDE_P2 = ("Day by day: what gets built, and the single test that proves each "
           "day is finished.")
META_P2 = [("Working days:", "5"),
           ("Blocking inputs:", "from you, Day 0"),
           ("Go-live:", "end of Day 5"),
           ("Hypercare:", "1 week after")]

TIMELINE_CAPTION = [
    ("Why this order:", {"bold": True, "color": SUB}),
    (" access first — nothing can be tested without it. Then the statutory path "
     "(Days 2–3), because an invoice has no legal validity without its IRN. "
     "Then the books (Day 4). Training and go-live last, so nobody is trained "
     "on something still moving.", {}),
]

# order is exactly as specified: DAY 1..5 then DAY 0 (row-major 2-col table)
DAYS = [
    {
        "tag": "DAY 1",
        "title": "Foundation & access",
        "bullets": [
            "Sandbox test credentials from the GSP partner configured and verified end to end.",
            "Your Tally chart of accounts collected; every ledger name mapped to the export.",
            "Existing invoice records audited for compliance data; a cleanup list prepared.",
        ],
        "footer_label": "DONE WHEN",
        "footer_text": ("the CRM talks to the sandbox, and the ledger map matches "
                        "your Tally company exactly."),
    },
    {
        "tag": "DAY 2",
        "title": "GSP integration — E-Invoice (IRN)",
        "bullets": [
            "Secure server-side connection built; credentials are never exposed to the browser.",
            "IRN generation, cancellation and details lookup implemented and tested against sandbox.",
            "Manual download-paste replaced by one button; kept as a fallback if the API is down.",
        ],
        "footer_label": "DONE WHEN",
        "footer_text": ("a test invoice receives a real IRN from the sandbox and "
                        "is marked Complete in the CRM."),
    },
    {
        "tag": "DAY 3",
        "title": "E-Way Bill & compliance gates",
        "bullets": [
            "E-Way Bill raised automatically after the IRN, with transport and dispatch details prefilled.",
            "24-hour cancellation window enforced in the interface; validity date stored on the invoice.",
            "Clear error messages, automatic retries, and a log of every API call.",
        ],
        "footer_label": "DONE WHEN",
        "footer_text": ("a test E-Way Bill is generated with a valid till date, "
                        "and cancellation after 24 hours is blocked."),
    },
    {
        "tag": "DAY 4",
        "title": "Tally export",
        "bullets": [
            "Export wired in: choose a date range and branch, download the Tally XML.",
            "Balance guard, one ledger per GST rate, and a drift report comparing CRM and Tally figures.",
            "Imported into a Tally trial company; voucher count verified.",
        ],
        "footer_label": "DONE WHEN",
        "footer_text": ("the Tally Sales register and Trial Balance tie back to "
                        "the CRM invoice register, to the paisa."),
    },
    {
        "tag": "DAY 5",
        "title": "End-to-end test, training & go-live",
        "bullets": [
            "Full run on real branch data: invoice → IRN → E-Way Bill → Tally.",
            "Production credentials swapped in and re-verified after the sandbox passes.",
            "Owner and accounts team trained; a one-page SOP handed over.",
        ],
        "footer_label": "DONE WHEN",
        "footer_text": ("the first real invoice completes the whole chain with "
                        "zero manual data entry."),
    },
    {
        "tag": "DAY 0",
        "title": "What we need from you first",
        "amber": True,
        "bullets": [
            [("Ledger names", {"bold": True}),
             (" from your TallyPrime company — Chart of Accounts → Ledgers.", {})],
            [("GSP sandbox credentials", {"bold": True}),
             (" and sign-off on the API proposal (4,000 calls/year tier).", {})],
            [("Accountant's call", {"bold": True}),
             (" on the one-paisa CGST/SGST rounding note.", {})],
        ],
        "footer_label": "Note",
        "footer_text": ("these three items are the only things that can delay "
                        "Day 1. Everything else is on our side."),
    },
]

ASSUMPTIONS = [
    "The CRM stays the only billing source — Tally consumes data; no two-way sync to maintain.",
    "One Tally company on one Windows PC; all three branches are supported from day one.",
    "Import starts file-based (no always-on gateway); the live gateway can be added later.",
    [("Not in these five days:", {"bold": True, "color": INK}),
     (" purchase and credit-note vouchers, payment allocation splitting, bulk "
      "backfill of old invoices.", {})],
]

COST_ROWS = [
    ("GSP APIs — e-Invoice + E-Way Bill + GSTIN, 4,000 calls/year",
     "₹29,500 / yr", False),
    ("On-boarding & support", "Free (FOC)", False),
    ("TallyPrime XML import gateway", "₹0 — built in", False),
    ("Hosting", "Existing — no change", False),
    ("Effective cost per e-invoiced despatch", "≈ ₹14.75", True),
]
COST_FOOTNOTE = ("An invoice that needs an E-Way Bill uses two calls (one IRN, "
                 "one E-Way Bill), so the 4,000-call tier covers roughly 2,000 "
                 "despatches a year.")

# ---------------------------------------------------------------------------
# layout constants for boxes (a spacer column keeps shaded boxes from touching)
# ---------------------------------------------------------------------------
BOX_W = 8.85
BOX_GAP = 0.30  # 8.85 + 0.30 + 8.85 = 18.0
STACK_W = [BOX_W, BOX_GAP, BOX_W]
SECTION5_W = [9.40, 0.40, 8.20]
LEGEND_W = [0.85, CONTENT_W - 0.85]
CHANGES_W = [3.60, 7.20, 7.20]


# ---------------------------------------------------------------------------
# building blocks
# ---------------------------------------------------------------------------
def masthead(flow, kicker, title, title_size, lede, lede_size, meta, space_after=5):
    table = add_table(flow._c, 1, 2, [11.4, CONTENT_W - 11.4])
    for cell in table.rows[0].cells:
        set_cell_margins(cell, bottom=0.22)
        set_cell_borders(cell, {"bottom": (NAVY, 18, "single")})  # 2.25 pt rule

    left = Flow(table.cell(0, 0))
    p = left.para(space_after=1, keep_with_next=True)
    rich(p, kicker, size=SZ_KICKER, bold=True, color=GREEN, spacing=15, caps=True)
    p = left.para(space_after=2, line_spacing=1.0, keep_with_next=True)
    rich(p, title, size=title_size, bold=True, color=NAVY, spacing=-5)
    p = left.para(space_after=0, line_spacing=1.12)
    rich(p, lede, size=lede_size, color=SUB)

    right = table.cell(0, 1)
    right.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    rflow = Flow(right)
    for index, (label, value) in enumerate(meta):
        last = index == len(meta) - 1
        p = rflow.para(align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.15,
                       space_after=0 if last else 1)
        rich(p, [(label, {"bold": True, "color": SUB}), (" " + value, {})],
             size=SZ_MICRO, color=MUTED)

    spacer = flow.para(space_after=space_after, line_spacing=1.0)
    style_run(spacer.add_run(""), size=2)
    return table


def section_heading(flow, number, title):
    p = flow.para(space_before=2, space_after=3.5, keep_with_next=True,
                  line_spacing=1.0)
    rich(p, " {} ".format(number), size=SZ_SMALL, bold=True, color=GREEN,
         shade=SOFT, spacing=4)
    rich(p, "  " + title, size=SZ_H2, bold=True, color=NAVY)
    return p


def card(flow, spec, fill=SOFT):
    p = flow.para(space_after=1.5, keep_with_next=True)
    rich(p, spec["tag"], size=SZ_MICRO, bold=True, color=GREEN, spacing=10)
    p = flow.para(space_after=2, line_spacing=1.1, keep_with_next=True)
    rich(p, spec["title"], size=SZ_CARD + 0.4, bold=True, color=NAVY)
    p = flow.para(space_after=4, line_spacing=1.12)
    rich(p, spec["body"], size=SZ_CARD, color=SUB)
    p = flow.para(space_after=0, line_spacing=1.0)
    rich(p, " " + spec["chip"] + " ", size=SZ_SMALL, bold=True, color=GREEN,
         shade=GREEN_TINT, spacing=2)


def build_page_one(doc):
    flow = Flow(doc, reuse_first=False)
    masthead(flow, "PROKON CRM · CONFIDENTIAL — FOR CLIENT REVIEW",
             "API Integration & Tally Export", SZ_H1, LEDE_P1, SZ_BODY, META_P1)

    # ---- 01 -----------------------------------------------------------------
    section_heading(flow, "01", "Two automations, one source of truth")
    table = add_table(doc, 1, 3, STACK_W)
    for cell in (table.cell(0, 0), table.cell(0, 2)):
        shade_cell(cell, SOFT)
        set_cell_margins(cell, top=0.20, bottom=0.20, left=0.24, right=0.24)
    card(Flow(table.cell(0, 0)), CARD_A)
    card(Flow(table.cell(0, 2)), CARD_B)

    p = flow.para(space_before=4, space_after=0, line_spacing=1.12,
                  left_indent=0.18, right_indent=0.18)
    shade_paragraph(p, GREEN_TINT)
    set_paragraph_borders(p, {"left": (GREEN, 19, "single", 6)})
    rich(p, RULE, size=SZ_CARD)

    # ---- 02 -----------------------------------------------------------------
    section_heading(flow, "02", "The process, end to end")
    p = flow.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3,
                  line_spacing=1.0, keep_with_next=True)
    p.add_run().add_picture(FIG_FLOW, width=Cm(IMAGE_W))

    legend = add_table(doc, len(LEGEND), 2, LEGEND_W)
    for index, (label, text) in enumerate(LEGEND):
        row = legend.rows[index]
        set_row_cant_split(row)
        left, right = row.cells
        set_cell_margins(left, right=0.10)
        set_cell_margins(right, left=0.10)
        lp = fmt_para(left.paragraphs[0], space_after=1.6, line_spacing=1.05)
        rich(lp, label, size=SZ_SMALL, bold=True, color=GREEN)
        rp = fmt_para(right.paragraphs[0], space_after=1.6, line_spacing=1.05)
        rich(rp, text, size=SZ_SMALL, color=SUB)

    # ---- 03 -----------------------------------------------------------------
    p = section_heading(flow, "03", "What changes for your team")
    p.paragraph_format.space_before = Pt(6)
    table = add_table(doc, 1 + len(CHANGES), 3, CHANGES_W)
    for index, text in enumerate(CHANGES_HEAD):
        cell = table.cell(0, index)
        set_cell_margins(cell, bottom=0.08)
        set_cell_borders(cell, {"bottom": (NAVY, 11, "single")})
        para = fmt_para(cell.paragraphs[0], space_after=2, line_spacing=1.05)
        rich(para, text.upper(), size=SZ_MICRO, bold=True, color=MUTED,
             spacing=8, caps=True)
    for r, row_data in enumerate(CHANGES, start=1):
        last = r == len(CHANGES)
        row = table.rows[r]
        set_row_cant_split(row)
        for c, text in enumerate(row_data):
            cell = row.cells[c]
            set_cell_margins(cell, top=0.10, bottom=0.10)
            if not last:
                set_cell_borders(cell, {"bottom": (BORDER, 6, "single")})
            para = fmt_para(cell.paragraphs[0], space_after=0, line_spacing=1.08)
            if c == 0:
                rich(para, text, size=SZ_TABLE, bold=True, color=INK)
            elif c == 1:
                rich(para, text, size=SZ_TABLE, color=MUTED)
            else:
                rich(para, text, size=SZ_TABLE, color=INK)

    doc.add_page_break()


def build_page_two(doc):
    flow = Flow(doc, reuse_first=False)
    masthead(flow, "PROKON CRM · DELIVERY PLAN", "Five-day setup plan",
             SZ_H1_P2, LEDE_P2, SZ_BODY, META_P2)

    # ---- timeline figure + caption -----------------------------------------
    p = flow.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3,
                  line_spacing=1.0, keep_with_next=True)
    p.add_run().add_picture(FIG_TIMELINE, width=Cm(IMAGE_W))
    p = flow.para(space_after=2, line_spacing=1.1)
    rich(p, TIMELINE_CAPTION, size=SZ_MICRO, color=MUTED)

    # ---- 04 -----------------------------------------------------------------
    section_heading(flow, "04", "Day by day")
    table = add_table(doc, len(DAYS), 3, [STACK_W[0], STACK_W[1], STACK_W[2]])
    for box_index, spec in enumerate(DAYS):
        row_index = box_index // 2
        header_cell = table.cell(row_index * 2, 0 if box_index % 2 == 0 else 2)
        body_cell = table.cell(row_index * 2 + 1, 0 if box_index % 2 == 0 else 2)
        amber = spec.get("amber", False)
        edge = AMBER_EDGE if amber else BORDER
        head_fill = AMBER if amber else NAVY
        tag_color = DAY_TAG_AMBER if amber else DAY_TAG
        body_fill = AMBER_TINT if amber else WHITE
        body_color = AMBER_BODY if amber else SUB

        set_row_cant_split(table.rows[row_index * 2])
        set_row_cant_split(table.rows[row_index * 2 + 1])

        # header bar ---------------------------------------------------------
        shade_cell(header_cell, head_fill)
        set_cell_margins(header_cell, top=0.12, bottom=0.10, left=0.22,
                         right=0.22)
        set_cell_borders(header_cell, {"top": (edge, 6, "single"),
                                       "left": (edge, 6, "single"),
                                       "right": (edge, 6, "single")})
        header_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        hp = fmt_para(header_cell.paragraphs[0], space_after=0,
                      line_spacing=1.0)
        rich(hp, spec["tag"], size=SZ_MICRO, bold=True, color=tag_color,
             spacing=12)
        rich(hp, "  " + spec["title"], size=SZ_TABLE, bold=True, color=WHITE)

        # body ---------------------------------------------------------------
        shade_cell(body_cell, body_fill)
        set_cell_margins(body_cell, top=0.12, bottom=0.14, left=0.22,
                         right=0.22)
        set_cell_borders(body_cell, {"left": (edge, 6, "single"),
                                     "right": (edge, 6, "single"),
                                     "bottom": (edge, 6, "single")})
        body_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        bflow = Flow(body_cell)
        for bullet in spec["bullets"]:
            bflow.bullet(bullet, size=SZ_BULLET, color=body_color,
                         space_after=1.4)
        foot = bflow.para(space_before=2, space_after=0, line_spacing=1.06)
        set_paragraph_borders(foot, {"top": (edge, 6, "single", 3)})
        rich(foot, spec["footer_label"], size=SZ_MICRO, bold=True,
             color=AMBER if amber else GREEN, spacing=8,
             font=FONT)
        rich(foot, " — " + spec["footer_text"], size=SZ_MICRO, color=body_color)

    # ---- 05 -----------------------------------------------------------------
    p = section_heading(flow, "05", "Assumptions, exclusions and cost")
    p.paragraph_format.space_before = Pt(6)
    table = add_table(doc, 1, 3, SECTION5_W)

    left_cell = table.cell(0, 0)
    set_cell_margins(left_cell, top=0.06, bottom=0.06, right=0.20)
    lflow = Flow(left_cell)
    for index, item in enumerate(ASSUMPTIONS):
        lflow.bullet(item, size=SZ_BULLET, color=SUB,
                     space_after=1.4 if index < len(ASSUMPTIONS) - 1 else 0)

    right_cell = table.cell(0, 2)
    shade_cell(right_cell, GREEN_TINT)
    set_cell_margins(right_cell, top=0.18, bottom=0.18, left=0.22, right=0.22)
    set_cell_borders(right_cell, {"top": (BORDER, 6, "single"),
                                  "bottom": (BORDER, 6, "single"),
                                  "left": (GREEN, 19, "single"),
                                  "right": (BORDER, 6, "single")})
    rflow = Flow(right_cell)
    p = rflow.para(space_after=2.5, line_spacing=1.05, keep_with_next=True)
    rich(p, "Cost — nothing new to buy for Tally", size=SZ_CARD, bold=True,
         color=GREEN)

    inner_w = SECTION5_W[2] - 0.44
    cost = add_table(right_cell, len(COST_ROWS), 2,
                     [round(inner_w * 0.62, 2), round(inner_w * 0.38, 2)])
    for index, (label, value, strong) in enumerate(COST_ROWS):
        row = cost.rows[index]
        set_row_cant_split(row)
        last = index == len(COST_ROWS) - 1
        lcell, vcell = row.cells
        for cell in (lcell, vcell):
            set_cell_margins(cell, top=0.06, bottom=0.06, left=0, right=0)
            if not last:
                set_cell_borders(cell, {"bottom": (BORDER, 6, "single")})
        lp = fmt_para(lcell.paragraphs[0], space_after=0, line_spacing=1.05)
        rich(lp, label, size=SZ_MICRO, bold=strong, color=COST_LABEL)
        vp = fmt_para(vcell.paragraphs[0], align=WD_ALIGN_PARAGRAPH.RIGHT,
                      space_after=0, line_spacing=1.05)
        rich(vp, value, size=SZ_MICRO, bold=strong, color=GREEN)

    p = rflow.para(space_before=3, space_after=0, line_spacing=1.1)
    rich(p, COST_FOOTNOTE, size=SZ_MICRO - 0.4, color=COST_LABEL)

    # a table may not be the last block in a body
    tail = flow.para(space_after=0, line_spacing=1.0)
    style_run(tail.add_run(""), size=2)


# ---------------------------------------------------------------------------
def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(SZ_BODY)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _insert_ordered(rPr, OxmlElement("w:rFonts"), RPR_ORDER)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.12
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)

    for style_name, size in (("List Bullet", SZ_BULLET),
                             ("Header", SZ_MICRO), ("Footer", SZ_MICRO)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(MUTED)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.space_before = Pt(0)


def configure_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(PAGE_W)
    section.page_height = Cm(PAGE_H)
    section.top_margin = Cm(MARGIN_TB)
    section.bottom_margin = Cm(MARGIN_TB)
    section.left_margin = Cm(MARGIN_LR)
    section.right_margin = Cm(MARGIN_LR)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    fmt_para(hp, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0,
             line_spacing=1.0)
    set_paragraph_borders(hp, {"bottom": (BORDER, 4, "single", 2)})
    rich(hp, "Prokon CRM · API Integration & Tally Export", size=SZ_MICRO,
         italic=True, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fmt_para(fp, space_after=0, line_spacing=1.0)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W),
                                               WD_TAB_ALIGNMENT.RIGHT)
    rich(fp, "Prokon CRM · API Integration & Tally Export · v1.0",
         size=SZ_MICRO, color=MUTED)
    rich(fp, "\tPage ", size=SZ_MICRO, color=MUTED)
    add_page_field(fp, size=SZ_MICRO, color=MUTED)


def main():
    for path in (FIG_FLOW, FIG_TIMELINE):
        if not os.path.isfile(path):
            sys.stderr.write("missing asset: {}\n".format(path))
            return 1

    doc = Document()
    # the default template ships one empty body paragraph; drop it so the
    # masthead table starts the document cleanly
    for para in list(doc.paragraphs):
        para._element.getparent().remove(para._element)

    configure_styles(doc)
    configure_section(doc)

    doc.core_properties.title = "API Integration & Tally Export"
    doc.core_properties.subject = ("Prokon CRM — e-Invoicing (IRN + E-Way Bill) "
                                   "and Tally export proposal")
    doc.core_properties.author = "Prokon CRM"

    build_page_one(doc)
    build_page_two(doc)

    out_dir = os.path.dirname(OUT_PATH)
    if not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    doc.save(OUT_PATH)

    size = os.path.getsize(OUT_PATH)
    print("wrote {} ({} bytes)".format(OUT_PATH, size))
    return 0


if __name__ == "__main__":
    sys.exit(main())
